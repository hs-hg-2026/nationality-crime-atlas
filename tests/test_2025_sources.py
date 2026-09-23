from pathlib import Path

from openpyxl import Workbook

from nationality_crime_atlas.census_2025 import (
    parse_census_estimated_foreign_population,
    parse_census_preliminary_population,
)
from nationality_crime_atlas.npa_2025_summary import (
    parse_npa_2025_foreign_clearance_totals,
    parse_npa_2025_prefecture_crime,
    parse_npa_2025_selected_nationalities,
    parse_npa_2025_visiting_offenses,
)


def _npa_prefecture_fixture(path: Path) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet, title, national, hokkaido, aomori in (
        ("1-5-1", "都道府県別刑法犯認知件数", 774_142, 24_150, 4_665),
        ("1-5-3", "都道府県別刑法犯検挙件数", 301_055, 12_633, 2_046),
        ("1-5-4", "都道府県別刑法犯検挙人員", 200_663, 9_431, 1_478),
    ):
        worksheet = workbook.create_sheet(sheet)
        worksheet["B1"] = f"図表：{sheet}（{title}）"
        worksheet["D2"] = "年次"
        worksheet["M2"] = "R6"
        worksheet["N2"] = "R7"
        worksheet["B4"] = "全国総数(件)"
        worksheet["N4"] = national
        worksheet["C5"] = "北海道"
        worksheet["N5"] = hokkaido
        worksheet["C6"] = "東北"
        worksheet["D6"] = "計"
        worksheet["N6"] = 99_999
        worksheet["D7"] = "青森"
        worksheet["N7"] = aomori
    workbook.save(path)
    return path


def _npa_foreign_fixture(path: Path) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)

    totals = workbook.create_sheet("3-3-1")
    totals["B1"] = "図表：３－３－１（外国人の刑法犯検挙状況）"
    totals["D2"] = "年次"
    totals["M2"] = "R6"
    totals["N2"] = "R7"
    for row, label, value in (
        (4, "検挙件数（件）", 301_055),
        (5, "うち外国人", 22_917),
        (6, "うち来日外国人", 17_614),
        (7, "検挙人員（人）", 200_663),
        (8, "うち外国人", 11_354),
        (9, "うち来日外国人", 7_333),
    ):
        totals.cell(row, 2 if row in {4, 7} else 3 if row in {5, 8} else 4).value = label
        totals.cell(row, 14).value = value

    nationality = workbook.create_sheet("3-3-3")
    nationality["B1"] = "図表：３－３－３（国籍等別刑法犯検挙状況）"
    nationality["D2"] = "年次"
    nationality["N2"] = "R7"
    rows = (
        (4, 2, "外国人検挙件数(件)", 22_917),
        (5, 4, "うち来日", 17_614),
        (6, 3, "中国", 3_321),
        (7, 4, "うち来日", 2_225),
        (8, 3, "タイ", 1_737),
        (9, 4, "うち来日", 1_687),
        (10, 2, "外国人検挙人員(人)", 11_354),
        (11, 4, "うち来日", 7_333),
        (12, 3, "中国", 2_343),
        (13, 4, "うち来日", 1_449),
        (14, 3, "ミャンマー", 171),
        (15, 4, "うち来日", 163),
    )
    for row, column, label, value in rows:
        nationality.cell(row, column).value = label
        nationality.cell(row, 14).value = value
    nationality["B16"] = (
        "※　検挙件数は300件以上の年がある国・地域を抽出\n"
        "※　検挙人員は150人以上の年がある国・地域を抽出\n"
        "※　中国に、「台湾」及び「香港等」は含まない。"
    )

    offenses = workbook.create_sheet("3-3-4")
    offenses["B1"] = "図表：３－３－４（来日外国人の罪種・手口別刑法犯検挙状況）"
    offenses["F2"] = "年次"
    offenses["P2"] = "R7"
    offenses["B4"] = "総数"
    offenses["F4"] = "検挙件数(件)"
    offenses["P4"] = 17_614
    offenses["F5"] = "検挙人員(人)"
    offenses["P5"] = 7_333
    offenses["C6"] = "凶悪犯"
    offenses["F6"] = "検挙件数"
    offenses["P6"] = 295
    offenses["F7"] = "検挙人員"
    offenses["P7"] = 341
    offenses["D8"] = "うち殺人"
    offenses["F8"] = "検挙件数"
    offenses["P8"] = 55
    offenses["F9"] = "検挙人員"
    offenses["P9"] = 62
    workbook.save(path)
    return path


def _census_total_fixture(path: Path) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "a01"
    worksheet["A1"] = "令和７年国勢調査　人口速報集計"
    worksheet["A2"] = "第１表　男女別人口－全国、都道府県、市区町村"
    worksheet["C9"] = "表章項目"
    worksheet["D9"] = "人口"
    worksheet["D11"] = "0_総数"
    worksheet["A13"] = "地域識別コード"
    worksheet["B13"] = "都道府県"
    worksheet["C13"] = "地域名"
    for row, values in enumerate(
        (
            ("a", "00_全国", "00000_全国", 123_049_524),
            ("a", "01_北海道", "01000_北海道", 4_985_419),
            ("1", "01_北海道", "01100_札幌市", 1_964_034),
            ("a", "02_青森県", "02000_青森県", 1_157_332),
        ),
        start=14,
    ):
        for column, value in enumerate(values, start=1):
            worksheet.cell(row, column).value = value
    workbook.save(path)
    return path


def _census_foreign_fixture(path: Path) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "［参考表］全国、都道府県、市区町村別外国人（推計）"
    worksheet["A3"] = "都道府県"
    worksheet["B3"] = "地域"
    worksheet["C3"] = "外国人(推計)"
    for row, values in enumerate(
        (
            ("00_全国", "00000_全国", 3_213_212),
            ("01_北海道", "01000_北海道", 58_246),
            ("01_北海道", "01100_札幌市", 18_343),
            ("02_青森県", "02000_青森県", 9_876),
        ),
        start=4,
    ):
        for column, value in enumerate(values, start=1):
            worksheet.cell(row, column).value = value
    workbook.save(path)
    return path


def test_npa_2025_prefecture_parser_joins_three_official_tables(tmp_path):
    records = parse_npa_2025_prefecture_crime(
        _npa_prefecture_fixture(tmp_path / "r07_1.xlsx"), source_id="S20"
    )

    assert len(records) == 3
    national = next(record for record in records if record.geography == "日本")
    aomori = next(record for record in records if record.geography == "青森県")
    assert (national.recognized_cases, national.cleared_cases, national.cleared_persons) == (
        774_142,
        301_055,
        200_663,
    )
    assert aomori.cleared_cases == 2_046
    assert aomori.geography_semantics == "police_reporting_area_unresolved"


def test_npa_2025_foreign_totals_preserve_three_population_scopes(tmp_path):
    records = parse_npa_2025_foreign_clearance_totals(
        _npa_foreign_fixture(tmp_path / "r07_3.xlsx"), source_id="S21"
    )

    assert len(records) == 3
    visiting = next(r for r in records if r.population_scope == "visiting_foreign")
    assert visiting.year == 2025
    assert (visiting.cleared_cases, visiting.cleared_persons) == (17_614, 7_333)


def test_npa_2025_selected_nationalities_keep_metric_specific_coverage(tmp_path):
    records = parse_npa_2025_selected_nationalities(
        _npa_foreign_fixture(tmp_path / "r07_3.xlsx"), source_id="S22"
    )

    assert len(records) == 8
    assert any(
        r.nationality == "タイ" and r.metric == "cleared_cases" for r in records
    )
    assert not any(
        r.nationality == "タイ" and r.metric == "cleared_persons" for r in records
    )
    assert any(
        r.nationality == "ミャンマー" and r.metric == "cleared_persons"
        for r in records
    )
    china = next(
        r
        for r in records
        if r.nationality == "中国"
        and r.metric == "cleared_cases"
        and r.population_scope == "all_foreign"
    )
    assert china.value == 3_321
    assert china.category_definition == "excludes_taiwan_and_hong_kong_etc"


def test_npa_2025_visiting_offenses_preserve_hierarchy(tmp_path):
    records = parse_npa_2025_visiting_offenses(
        _npa_foreign_fixture(tmp_path / "r07_3.xlsx"), source_id="S23"
    )

    assert len(records) == 3
    heinous = next(record for record in records if record.offense_label == "凶悪犯")
    murder = next(record for record in records if record.offense_label == "殺人")
    assert (heinous.cleared_cases, heinous.cleared_persons) == (295, 341)
    assert murder.offense_parent_id == heinous.offense_id


def test_census_preliminary_population_keeps_only_national_and_prefecture_rows(tmp_path):
    records = parse_census_preliminary_population(
        _census_total_fixture(tmp_path / "census.xlsx"), source_id="S24"
    )

    assert [record.geography for record in records] == ["日本", "北海道", "青森県"]
    assert records[0].population == 123_049_524
    assert records[0].rounding == "none"
    assert records[0].geography_semantics == "census_preliminary_population"


def test_census_estimated_foreign_population_is_explicitly_estimated(tmp_path):
    records = parse_census_estimated_foreign_population(
        _census_foreign_fixture(tmp_path / "foreign.xlsx"), source_id="S25"
    )

    assert [record.geography for record in records] == ["日本", "北海道", "青森県"]
    assert records[0].population == 3_213_212
    assert records[0].population_scope == "foreign_population_estimate"
    assert records[0].geography_semantics == "census_reference_estimate"
