from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from nationality_crime_atlas.census_2025 import (
    parse_census_final_total_population,
    parse_census_final_nationality_population,
    parse_census_final_detailed_nationalities,
)
from nationality_crime_atlas.errors import SchemaError


def _total_fixture(path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "b01_01"
    sheet["A1"] = "【原数値】令和７年国勢調査　人口等基本集計"
    sheet["A2"] = "第１－１表　男女別人口"
    sheet["H10"] = "人口"
    sheet["H12"] = "0_総数"
    sheet["H14"] = "人"
    for row, values in enumerate([
        ("a", "00_全国", "00000", "2000", "00_全国", "00000", "0001_全国", 1000),
        ("a", "01_北海道", "01000", "2000", "01_北海道", "01000", "0002_北海道", 1000),
        ("1", "01_北海道", "01100", "2000", "01_北海道", "01100", "0003_札幌市", 800),
        ("9", "01_北海道", "01999", "2000", "01_北海道", "01000", "0004_（旧：例）", 10),
    ], start=16):
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column, value)
    workbook.save(path)
    return path


def _nationality_fixture(path: Path, *, imputed=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "b49_01"
    sheet["A1"] = ("【不詳補完値】" if imputed else "【原数値】") + "令和７年国勢調査　人口等基本集計"
    sheet["A2"] = "第４９－１表　男女、国籍別人口"
    categories = ["0_総数", "1_外国人", "102_中国", "124_その他", "2_日本人", "3_日本人・外国人の別「不詳」"]
    for column, category in enumerate(categories, start=5):
        sheet.cell(5, column, "人口")
        sheet.cell(7, column, category)
        sheet.cell(8, column, 2 if category.startswith(("102_", "124_")) else 1)
        sheet.cell(9, column, "人")
    for row, values in enumerate([
        ("0_総数", "a", "00_全国", "00000_全国", 1000, 200, 150, 50, 790, 10),
        ("0_総数", "a", "01_北海道", "01000_北海道", 1000, 200, 150, 50, 790, 10),
        ("0_総数", "1", "01_北海道", "01100_札幌市", 500, 100, 80, 20, 400, "-"),
        ("1_男", "a", "00_全国", "00000_全国", 600, 120, 90, 30, 480, "-"),
    ], start=11):
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column, value)
    workbook.save(path)
    return path


def _detailed_fixture(path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "b55"
    sheet["A1"] = "【原数値】令和７年国勢調査　人口等基本集計"
    sheet["A2"] = "第５５表　男女、国籍（詳細区分）別人口－全国"
    sheet["D5"] = "人口"
    sheet["D7"] = "0_総数"
    sheet["D9"] = "人"
    for row, values in enumerate([
        ("00000_全国", "1", "0_総数", 1000),
        ("00000_全国", "1", "1_外国人", 200),
        ("00000_全国", "2", "11_アジア州", 150),
        ("00000_全国", "3", "1101_アゼルバイジャン", 150),
        ("00000_全国", "2", "17_無国籍・国名「不詳」", 50),
        ("00000_全国", "1", "2_日本人", 790),
        ("00000_全国", "1", "3_日本人・外国人の別「不詳」", 10),
    ], start=11):
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column, value)
    workbook.save(path)
    return path


def _change(path, cell, value):
    workbook = load_workbook(path)
    workbook.active[cell] = value
    workbook.save(path)


def test_final_total_excludes_cities_and_old_boundaries(tmp_path):
    rows = parse_census_final_total_population(_total_fixture(tmp_path / "total.xlsx"), source_id="S26")
    assert [(r.geography, r.population) for r in rows] == [("日本", 1000), ("北海道", 1000)]
    assert rows[0].geography_semantics == "census_final_original_population"
    assert rows[0].rounding == "none"


def test_nationality_keeps_japanese_foreign_unknown_and_hierarchy(tmp_path):
    rows = parse_census_final_nationality_population(_nationality_fixture(tmp_path / "nat.xlsx"), source_id="S27")
    assert len(rows) == 12
    national = {r.nationality_code: r for r in rows if r.geography == "日本"}
    assert national["2"].population == 790
    assert national["3"].population == 10
    assert national["0"].population != national["1"].population + national["2"].population
    assert national["124"].nationality == "その他"
    assert national["124"].category_level == 2
    assert national["2"].value_basis == "original"
    assert national["2"].source_column == 9


def test_imputed_basis_and_source_dash_are_preserved(tmp_path):
    path = _nationality_fixture(tmp_path / "imputed.xlsx", imputed=True)
    _change(path, "J12", "-")
    _change(path, "I12", 800)
    rows = parse_census_final_nationality_population(path, source_id="S29", imputed=True)
    unknown = next(r for r in rows if r.geography == "北海道" and r.nationality_code == "3")
    assert unknown.population == 0
    assert unknown.source_value == "-"
    assert unknown.value_basis == "official_imputed_reference"


@pytest.mark.parametrize("imputed", [False, True])
def test_rejects_mislabeled_value_basis(tmp_path, imputed):
    path = _nationality_fixture(tmp_path / "nat.xlsx", imputed=not imputed)
    with pytest.raises(SchemaError, match="basis"):
        parse_census_final_nationality_population(path, source_id="S27", imputed=imputed)


@pytest.mark.parametrize("cell,value", [("I11", 789), ("G11", 149), ("H11", None), ("H11", "X"), ("G7", "2_日本人"), ("D12", "02000_北海道")])
def test_rejects_invalid_counts_headers_or_geography(tmp_path, cell, value):
    path = _nationality_fixture(tmp_path / "nat.xlsx")
    _change(path, cell, value)
    with pytest.raises(SchemaError):
        parse_census_final_nationality_population(path, source_id="S27")


def test_detail_preserves_region_and_stateless_rows(tmp_path):
    rows = parse_census_final_detailed_nationalities(_detailed_fixture(tmp_path / "detail.xlsx"), source_id="S28")
    assert len(rows) == 7
    stateless = next(r for r in rows if r.nationality_code == "17")
    assert stateless.nationality == "無国籍・国名「不詳」"
    assert stateless.category_level == 2
    assert stateless.geography == "日本"


def test_detail_rejects_inconsistent_region_total(tmp_path):
    path = _detailed_fixture(tmp_path / "detail.xlsx")
    _change(path, "D14", 149)
    with pytest.raises(SchemaError, match="reconcile"):
        parse_census_final_detailed_nationalities(path, source_id="S28")


@pytest.mark.parametrize("parser,source_id,record_type,count", [
    ("census-2025-final-total-population", "S26", "prefecture_population", 2),
    ("census-2025-final-nationality-population", "S27", "census_nationality_population", 12),
    ("census-2025-final-detailed-nationalities", "S28", "census_nationality_population", 7),
    ("census-2025-imputed-nationality-population", "S29", "census_nationality_population", 12),
])
def test_final_census_runs_through_quality_pipeline_and_reuses(tmp_path, parser, source_id, record_type, count):
    import hashlib
    import json
    from nationality_crime_atlas.pipeline import run_offline_pipeline
    from nationality_crime_atlas.registry import load_source_registry

    source = load_source_registry("config/sources.json")[source_id]
    assert source["parser"] == parser
    if source_id == "S26":
        path = _total_fixture(tmp_path / "fixture.xlsx")
    elif source_id == "S28":
        path = _detailed_fixture(tmp_path / "fixture.xlsx")
    else:
        path = _nationality_fixture(tmp_path / "fixture.xlsx", imputed=source_id == "S29")
    # Pin only the synthetic fixture in this isolated pipeline; keep official pins intact.
    source["expected_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    profile = {"record_type": record_type, "expected_record_count": count, "expected_years": [2025]}
    if source_id != "S26":
        profile["allowed_values"] = {"value_basis": ["official_imputed_reference" if source_id == "S29" else "original"]}
    options = dict(source_id=source_id, source_metadata=source, quality_profile=profile,
                   retrieved_at="2026-10-05T21:00:00+09:00", raw_root=tmp_path / "raw",
                   processed_root=tmp_path / "processed")
    first = run_offline_pipeline(path, **options)
    assert json.loads(first.quality_report_path.read_text())["passed"]
    assert run_offline_pipeline(path, **options).reused
