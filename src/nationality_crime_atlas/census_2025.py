"""Parsers for the 2025 Population Census preliminary reference tables."""

from pathlib import Path
from typing import List, Optional, Tuple

from openpyxl import load_workbook

from .errors import SchemaError
from .models import CensusNationalityPopulationRecord, PrefecturePopulationRecord
from .npa_all_residents import PREFECTURE_BASE_LABELS, PREFECTURE_PARENT_REGION, _canonical_prefecture


def _split_code_label(value: object) -> Tuple[str, str]:
    text = str(value or "").strip()
    if "_" not in text:
        raise SchemaError("Census geography must contain a code and label: %r" % text)
    code, label = text.split("_", 1)
    return code, label.strip()


def _integer(value: object, *, row: int) -> int:
    try:
        numeric = float(value)
    except (TypeError, ValueError) as error:
        raise SchemaError("Missing census population at row %d" % row) from error
    if not numeric.is_integer() or numeric < 0:
        raise SchemaError("Invalid census population at row %d" % row)
    return int(numeric)


def _population_record(
    *,
    geography: str,
    population: int,
    scope: str,
    semantics: str,
    source_id: str,
    source_table: str,
    source_sheet: str,
    source_row: int,
) -> PrefecturePopulationRecord:
    national = geography == "日本"
    return PrefecturePopulationRecord(
        year=2025,
        reference_date="2025-10-01",
        population_scope=scope,
        geography=geography,
        geography_type="national" if national else "prefecture",
        parent_region=None if national else PREFECTURE_PARENT_REGION[geography],
        geography_semantics=semantics,
        population=population,
        source_value=population,
        source_unit="persons",
        rounding="none",
        source_id=source_id,
        source_table=source_table,
        source_sheet=source_sheet,
        source_row=source_row,
    )


def _canonical_census_geography(code: str, label: str) -> Optional[str]:
    if code == "00000" and label == "全国":
        return "日本"
    if len(code) == 5 and code.endswith("000") and label in PREFECTURE_PARENT_REGION:
        return label
    return None


def parse_census_preliminary_population(
    path: Path, *, source_id: str
) -> List[PrefecturePopulationRecord]:
    """Parse exact-person preliminary total population for Japan and 47 prefectures."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        title = " ".join(
            str(worksheet.cell(row, 1).value or "") for row in (1, 2)
        )
        if "令和７年国勢調査" not in title or "人口速報集計" not in title:
            raise SchemaError("2025 census preliminary population title was not recognized")
        records = []
        for row, values in enumerate(
            worksheet.iter_rows(min_row=14, values_only=True), start=14
        ):
            code, label = _split_code_label(values[2])
            geography = _canonical_census_geography(code, label)
            if geography is None:
                continue
            records.append(
                _population_record(
                    geography=geography,
                    population=_integer(values[3], row=row),
                    scope="total_population",
                    semantics="census_preliminary_population",
                    source_id=source_id,
                    source_table="1",
                    source_sheet=worksheet.title,
                    source_row=row,
                )
            )
        return records
    finally:
        workbook.close()


def parse_census_estimated_foreign_population(
    path: Path, *, source_id: str
) -> List[PrefecturePopulationRecord]:
    """Parse the explicitly estimated foreign population reference table."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        if "外国人（推計）" not in str(worksheet.cell(1, 1).value or ""):
            raise SchemaError("2025 census estimated foreign-population title was not recognized")
        records = []
        for row, values in enumerate(
            worksheet.iter_rows(min_row=4, values_only=True), start=4
        ):
            code, label = _split_code_label(values[1])
            geography = _canonical_census_geography(code, label)
            if geography is None:
                continue
            records.append(
                _population_record(
                    geography=geography,
                    population=_integer(values[2], row=row),
                    scope="foreign_population_estimate",
                    semantics="census_reference_estimate",
                    source_id=source_id,
                    source_table="reference_foreign_population",
                    source_sheet=worksheet.title,
                    source_row=row,
                )
            )
        return records
    finally:
        workbook.close()


def _final_title(worksheet, table_title: str, *, imputed: bool = False) -> None:
    title = str(worksheet.cell(1, 1).value or "")
    basis = "【不詳補完値】" if imputed else "【原数値】"
    if not title.startswith(basis):
        raise SchemaError("Census value basis differs from the registered source")
    if "令和７年国勢調査" not in title or "人口等基本集計" not in title:
        raise SchemaError("Final 2025 census title was not recognized")
    if not str(worksheet.cell(2, 1).value or "").startswith(table_title):
        raise SchemaError("Final census table identity was not recognized")


def _final_geography(code: str, label: str) -> Optional[str]:
    expected = {"00000": "全国"}
    expected.update({
        "%02d000" % index: _canonical_prefecture(base)
        for index, base in enumerate(PREFECTURE_BASE_LABELS, start=1)
    })
    if code not in expected:
        return None
    if expected[code] != label:
        raise SchemaError("Census geography code and label disagree: %s %s" % (code, label))
    return "日本" if code == "00000" else label


def _require_cells(worksheet, expected) -> None:
    for cell, value in expected.items():
        if worksheet[cell].value != value:
            raise SchemaError("Census header differs at %s" % cell)


def parse_census_final_total_population(
    path: Path, *, source_id: str
) -> List[PrefecturePopulationRecord]:
    """Read Table 1-1 current-boundary totals; exclude cities and historical rows."""
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        _final_title(sheet, "第１－１表")
        _require_cells(sheet, {"H10": "人口", "H12": "0_総数", "H14": "人"})
        records = []
        seen = set()
        for row, values in enumerate(sheet.iter_rows(min_row=16, values_only=True), start=16):
            if values[0] != "a":
                continue
            _, label = _split_code_label(values[6])
            code = str(values[5])
            geography = _final_geography(code, label)
            if geography is None or code in seen:
                raise SchemaError("Unexpected or duplicate census geography")
            seen.add(code)
            records.append(_population_record(
                geography=geography, population=_integer(values[7], row=row),
                scope="total_population", semantics="census_final_original_population",
                source_id=source_id, source_table="1-1", source_sheet=sheet.title, source_row=row,
            ))
        if not records:
            raise SchemaError("No final census population rows")
        return records
    finally:
        workbook.close()


def _census_category_record(*, code, geography, category, level, value, imputed,
                            source_id, table, sheet, row, column):
    category_code, label = _split_code_label(category)
    # The official user guide defines '-' as no applicable count. Preserve the symbol.
    population = 0 if value == "-" else _integer(value, row=row)
    return CensusNationalityPopulationRecord(
        year=2025, reference_date="2025-10-01", geography=geography,
        geography_code=code, geography_type="national" if code == "00000" else "prefecture",
        nationality_code=category_code, nationality=label, category_level=level,
        population=population, source_value="-" if value == "-" else population,
        source_unit="persons", rounding="none",
        value_basis="official_imputed_reference" if imputed else "original",
        source_id=source_id, source_table=table, source_sheet=sheet, source_row=row,
        source_column=column,
    )


def _reconcile_census_categories(records) -> None:
    for geography in {r.geography for r in records}:
        rows = [r for r in records if r.geography == geography]
        by_code = {r.nationality_code: r for r in rows}
        if len(by_code) != len(rows) or not {"0", "1", "2", "3"}.issubset(by_code):
            raise SchemaError("Duplicate or missing census nationality categories")
        if by_code["0"].population != sum(by_code[c].population for c in ("1", "2", "3")):
            raise SchemaError("Census total/Japanese/foreign/unknown counts do not reconcile")
        if by_code["1"].population != sum(r.population for r in rows if r.category_level == 2):
            raise SchemaError("Census foreign-category counts do not reconcile")
        parent = None
        children = []
        for record in rows + [None]:
            if record is None or record.category_level != 3:
                if children and sum(r.population for r in children) != parent.population:
                    raise SchemaError("Census region/country counts do not reconcile")
                parent = record if record is not None and record.category_level == 2 else None
                children = []
            else:
                if parent is None:
                    raise SchemaError("Census country has no region parent")
                children.append(record)


def parse_census_final_nationality_population(
    path: Path, *, source_id: str, imputed: bool = False
) -> List[CensusNationalityPopulationRecord]:
    """Read Table 49-1, retaining all nationality categories for 48 geographies."""
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        _final_title(sheet, "第４９－１表", imputed=imputed)
        columns = []
        for column in range(5, sheet.max_column + 1):
            category = sheet.cell(7, column).value
            code, _ = _split_code_label(category)
            level = _integer(sheet.cell(8, column).value, row=8)
            if level != (1 if code in {"0", "1", "2", "3"} else 2):
                raise SchemaError("Census nationality hierarchy changed")
            if sheet.cell(5, column).value != "人口" or sheet.cell(9, column).value != "人":
                raise SchemaError("Census nationality metric/unit changed")
            columns.append((column, category, level))
        records = []
        seen = set()
        for row, values in enumerate(sheet.iter_rows(min_row=11, values_only=True), start=11):
            if values[0] != "0_総数" or values[1] != "a":
                continue
            code, label = _split_code_label(values[3])
            geography = _final_geography(code, label)
            if geography is None or code in seen:
                raise SchemaError("Unexpected or duplicate census geography")
            seen.add(code)
            for column, category, level in columns:
                records.append(_census_category_record(
                    code=code, geography=geography, category=category, level=level,
                    value=values[column - 1], imputed=imputed, source_id=source_id,
                    table="49-1", sheet=sheet.title, row=row, column=column,
                ))
        if not records:
            raise SchemaError("No final census nationality rows")
        _reconcile_census_categories(records)
        return records
    finally:
        workbook.close()


def parse_census_final_detailed_nationalities(
    path: Path, *, source_id: str
) -> List[CensusNationalityPopulationRecord]:
    """Read Table 55 national detail, keeping country, region and unknown labels."""
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        _final_title(sheet, "第５５表")
        _require_cells(sheet, {"D5": "人口", "D7": "0_総数", "D9": "人"})
        records = []
        for row, values in enumerate(sheet.iter_rows(min_row=11, values_only=True), start=11):
            code, label = _split_code_label(values[0])
            if code != "00000" or label != "全国":
                raise SchemaError("Census detail must contain national rows only")
            level = _integer(values[1], row=row)
            if level not in {1, 2, 3}:
                raise SchemaError("Census detail hierarchy changed")
            records.append(_census_category_record(
                code=code, geography="日本", category=values[2], level=level, value=values[3],
                imputed=False, source_id=source_id, table="55", sheet=sheet.title, row=row, column=4,
            ))
        _reconcile_census_categories(records)
        return records
    finally:
        workbook.close()
