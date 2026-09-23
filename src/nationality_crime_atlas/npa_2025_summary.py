"""Parsers for the NPA 2025 criminal-code statistical summary workbooks."""

from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from .errors import SchemaError
from .models import (
    NationalClearanceAnnualRecord,
    NationalityOffenseGroupRecord,
    OverallPrefectureCrimeRecord,
    SelectedNationalityClearanceRecord,
)
from .npa_all_residents import PREFECTURE_PARENT_REGION, _canonical_prefecture


def _clean(value: object) -> str:
    return "".join(str(value or "").split())


def _integer(value: object, *, sheet: str, row: int) -> int:
    try:
        numeric = float(value)
    except (TypeError, ValueError) as error:
        raise SchemaError("Missing numeric value at %s row %d" % (sheet, row)) from error
    if not numeric.is_integer() or numeric < 0:
        raise SchemaError("Invalid numeric value at %s row %d" % (sheet, row))
    return int(numeric)


def _r7_column(worksheet, *, header_row: int = 2) -> int:
    matches = [
        column
        for column in range(1, worksheet.max_column + 1)
        if _clean(worksheet.cell(header_row, column).value) in {"R7", "令和7年"}
    ]
    if len(matches) != 1:
        raise SchemaError("Exactly one R7 column is required in sheet %s" % worksheet.title)
    return matches[0]


def _require_sheet(workbook, name: str, title_token: str):
    if name not in workbook.sheetnames:
        raise SchemaError("Required sheet %s was not found" % name)
    worksheet = workbook[name]
    if title_token not in _clean(worksheet.cell(1, 2).value):
        raise SchemaError("Sheet %s title was not recognized" % name)
    return worksheet


def _prefecture_metric(worksheet) -> Dict[str, Tuple[int, int]]:
    column = _r7_column(worksheet)
    values: Dict[str, Tuple[int, int]] = {}
    for row in range(4, worksheet.max_row + 1):
        col2 = _clean(worksheet.cell(row, 2).value)
        col3 = _clean(worksheet.cell(row, 3).value)
        col4 = _clean(worksheet.cell(row, 4).value)
        if col2.startswith("全国総数"):
            geography = "日本"
        elif col4 == "計":
            continue
        else:
            label = col4 or col3
            geography = _canonical_prefecture(label) if label else None
            if geography is None:
                continue
        if geography in values:
            raise SchemaError("Duplicate geography %s in sheet %s" % (geography, worksheet.title))
        values[geography] = (
            _integer(worksheet.cell(row, column).value, sheet=worksheet.title, row=row),
            row,
        )
    if "日本" not in values:
        raise SchemaError("National total was not found in sheet %s" % worksheet.title)
    return values


def parse_npa_2025_prefecture_crime(
    path: Path, *, source_id: str
) -> List[OverallPrefectureCrimeRecord]:
    """Join 2025 recognized cases, cleared cases, and persons by prefecture."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        recognized_sheet = _require_sheet(workbook, "1-5-1", "都道府県別刑法犯認知件数")
        cases_sheet = _require_sheet(workbook, "1-5-3", "都道府県別刑法犯検挙件数")
        persons_sheet = _require_sheet(workbook, "1-5-4", "都道府県別刑法犯検挙人員")
        recognized = _prefecture_metric(recognized_sheet)
        cases = _prefecture_metric(cases_sheet)
        persons = _prefecture_metric(persons_sheet)
        if set(recognized) != set(cases) or set(recognized) != set(persons):
            raise SchemaError("The three prefecture tables do not have identical geographies")
        records = []
        for geography, (recognized_value, source_row) in recognized.items():
            national = geography == "日本"
            records.append(
                OverallPrefectureCrimeRecord(
                    year=2025,
                    population_scope="all_persons",
                    offense_scope="criminal_code_excluding_traffic_negligence",
                    geography=geography,
                    geography_type="national" if national else "prefecture",
                    parent_region=None if national else PREFECTURE_PARENT_REGION[geography],
                    geography_semantics=(
                        "national_aggregate"
                        if national
                        else "police_reporting_area_unresolved"
                    ),
                    recognized_cases=recognized_value,
                    cleared_cases=cases[geography][0],
                    cleared_persons=persons[geography][0],
                    source_id=source_id,
                    source_table="1-5-1|1-5-3|1-5-4",
                    source_sheet="1-5-1|1-5-3|1-5-4",
                    source_row=source_row,
                )
            )
        return records
    finally:
        workbook.close()


def parse_npa_2025_foreign_clearance_totals(
    path: Path, *, source_id: str
) -> List[NationalClearanceAnnualRecord]:
    """Parse the three 2025 national population scopes from summary table 3-3-1."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = _require_sheet(workbook, "3-3-1", "外国人の刑法犯検挙状況")
        column = _r7_column(worksheet)
        specs = (
            ("all_persons", 4, 7),
            ("all_foreign", 5, 8),
            ("visiting_foreign", 6, 9),
        )
        return [
            NationalClearanceAnnualRecord(
                year=2025,
                population_scope=scope,
                offense_scope="criminal_code_excluding_traffic_negligence",
                geography="日本全国",
                cleared_cases=_integer(
                    worksheet.cell(cases_row, column).value,
                    sheet=worksheet.title,
                    row=cases_row,
                ),
                cleared_persons=_integer(
                    worksheet.cell(persons_row, column).value,
                    sheet=worksheet.title,
                    row=persons_row,
                ),
                source_id=source_id,
                source_table="3-3-1",
                source_sheet=worksheet.title,
                source_row=cases_row,
                source_cases_column=column,
                source_persons_column=column,
            )
            for scope, cases_row, persons_row in specs
        ]
    finally:
        workbook.close()


def parse_npa_2025_selected_nationalities(
    path: Path, *, source_id: str
) -> List[SelectedNationalityClearanceRecord]:
    """Parse only nationality metrics explicitly selected for publication in 3-3-3."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = _require_sheet(workbook, "3-3-3", "国籍等別刑法犯検挙状況")
        column = _r7_column(worksheet)
        records = []
        metric: Optional[str] = None
        current_nationality: Optional[str] = None
        for row in range(4, worksheet.max_row + 1):
            primary = _clean(worksheet.cell(row, 2).value)
            nationality = _clean(worksheet.cell(row, 3).value)
            detail = _clean(worksheet.cell(row, 4).value)
            if primary.startswith("※"):
                break
            if "外国人検挙件数" in primary:
                metric = "cleared_cases"
                current_nationality = None
                continue
            if "外国人検挙人員" in primary:
                metric = "cleared_persons"
                current_nationality = None
                continue
            if metric is None:
                continue
            if nationality:
                current_nationality = nationality
                population_scope = "all_foreign"
            elif detail == "うち来日" and current_nationality:
                population_scope = "visiting_foreign"
            else:
                continue
            records.append(
                SelectedNationalityClearanceRecord(
                    year=2025,
                    population_scope=population_scope,
                    nationality=current_nationality,
                    metric=metric,
                    value=_integer(
                        worksheet.cell(row, column).value,
                        sheet=worksheet.title,
                        row=row,
                    ),
                    selection_rule=(
                        "at_least_300_cases_in_one_year_2016_2025"
                        if metric == "cleared_cases"
                        else "at_least_150_persons_in_one_year_2016_2025"
                    ),
                    category_definition=(
                        "excludes_taiwan_and_hong_kong_etc"
                        if current_nationality == "中国"
                        else "as_published"
                    ),
                    source_id=source_id,
                    source_table="3-3-3",
                    source_sheet=worksheet.title,
                    source_row=row,
                    source_column=column,
                )
            )
        return records
    finally:
        workbook.close()


OFFENSE_IDS = {
    "総数": "criminal_code",
    "凶悪犯": "heinous",
    "殺人": "murder",
    "強盗": "robbery",
    "侵入強盗": "residential_or_building_robbery",
    "非侵入強盗": "non_intrusion_robbery",
    "粗暴犯": "assaultive",
    "傷害": "bodily_injury",
    "窃盗犯": "theft",
    "侵入窃盗": "intrusion_theft",
    "住宅対象": "residential_intrusion_theft",
    "乗り物盗": "vehicle_theft",
    "自動車盗": "motor_vehicle_theft",
    "非侵入窃盗": "non_intrusion_theft",
    "部品ねらい": "vehicle_parts_theft",
    "車上ねらい": "theft_from_vehicle",
    "ひったくり": "snatch_theft",
    "すり": "pickpocketing",
    "自動販売機ねらい": "vending_machine_theft",
    "万引き": "shoplifting",
    "知能犯": "intellectual",
    "偽造": "forgery",
    "風俗犯": "morals",
    "不同意わいせつ": "nonconsensual_indecency",
    "その他の刑法犯": "other_criminal_code",
    "占有離脱物横領": "embezzlement_of_lost_property",
    "住居侵入": "trespass",
    "略取誘拐・人身売買": "kidnapping_and_human_trafficking",
}


def _offense_label(worksheet, row: int) -> Tuple[Optional[str], Optional[int]]:
    for column in range(2, 6):
        label = _clean(worksheet.cell(row, column).value)
        if label:
            return label.removeprefix("うち"), column - 2
    return None, None


def parse_npa_2025_visiting_offenses(
    path: Path, *, source_id: str
) -> List[NationalityOffenseGroupRecord]:
    """Parse the 2025 visiting-foreigner offense hierarchy from table 3-3-4."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = _require_sheet(workbook, "3-3-4", "来日外国人の罪種・手口別刑法犯検挙状況")
        column = _r7_column(worksheet)
        records = []
        parents: Dict[int, str] = {}
        row = 4
        while row <= worksheet.max_row:
            label, level = _offense_label(worksheet, row)
            metric = _clean(worksheet.cell(row, 6).value)
            if not label or "検挙件数" not in metric:
                row += 1
                continue
            offense_id = OFFENSE_IDS.get(label)
            if offense_id is None:
                raise SchemaError("Unrecognized 3-3-4 offense label: %s" % label)
            persons_row = row + 1
            if "検挙人員" not in _clean(worksheet.cell(persons_row, 6).value):
                raise SchemaError("Missing paired persons row after %s" % label)
            parent_id = parents.get(level - 1) if level and level > 0 else None
            parents[level or 0] = offense_id
            records.append(
                NationalityOffenseGroupRecord(
                    year=2025,
                    population_scope="visiting_foreign",
                    region=None,
                    nationality=None,
                    subcategory=None,
                    row_kind="national_total",
                    offense_id=offense_id,
                    offense_label=label,
                    offense_parent_id=parent_id,
                    offense_level=level or 0,
                    official_severity_role=(
                        "official_high_severity_category"
                        if offense_id == "heinous"
                        else "not_a_project_severity_classification"
                    ),
                    cleared_cases=_integer(
                        worksheet.cell(row, column).value,
                        sheet=worksheet.title,
                        row=row,
                    ),
                    cleared_persons=_integer(
                        worksheet.cell(persons_row, column).value,
                        sheet=worksheet.title,
                        row=persons_row,
                    ),
                    source_id=source_id,
                    source_table="3-3-4",
                    source_sheet=worksheet.title,
                    source_row=row,
                    source_cases_column=column,
                    source_persons_column=column,
                )
            )
            row += 2
        return records
    finally:
        workbook.close()
